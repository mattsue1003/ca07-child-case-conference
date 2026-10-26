# 由 pptx_build/layout.json 組成「可編輯」的 PowerPoint：文字框、卡片、影片都是原生物件
# 用法：node layout.cjs && python3 build_pptx.py
import json, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

PX = 9525                      # 1280px 寬 = 13.333in
FONT_BODY = 'Microsoft JhengHei'
FONT_HEAD = 'Microsoft JhengHei'
BG = RGBColor(0xFB, 0xF3, 0xE6)
TERRA = RGBColor(0xD9, 0x73, 0x4E)
NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'

def E(v): return Emu(int(round(v * PX)))
def rgb(s):
    m = re.findall(r'[\d.]+', s or '')
    if len(m) < 3: return None
    if len(m) == 4 and float(m[3]) == 0: return None
    return RGBColor(*[int(float(x)) for x in m[:3]])

def set_font(run, name, size_px, bold, color):
    f = run.font
    f.name = name; f.size = Pt(size_px * 0.75); f.bold = bold
    if color: f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None: el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', name)

def alpha(shape, pct):
    sf = shape.fill._xPr.find(qn('a:solidFill'))
    clr = sf[0]
    a = etree.SubElement(clr, qn('a:alpha')); a.set('val', str(int(pct * 1000)))

def rrect(slide, x, y, w, h, fill, radius=0, line=None, lw=0, dash=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    if radius: shp.adjustments[0] = min(0.5, radius / max(1, min(w, h)))
    if fill: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    else: shp.fill.background()
    if line and lw:
        shp.line.color.rgb = line; shp.line.width = Pt(lw * 0.75)
        if dash: shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else: shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ''
    return shp

def timing(spid):
    return f'''<p:timing xmlns:p="{NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:par><p:cTn id="2" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="5" dur="indefinite" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
<p:video><p:cMediaNode vol="80000" mute="1" numSld="999" showWhenStopped="1"><p:cTn id="6" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''

def textbox(slide, t):
    x, y, w, h = t['x'], t['y'], t['w'], t['h']
    sz = t['sz']
    if t['li']:
        ind = sz * 1.15; x -= ind; w += ind
    tb = slide.shapes.add_textbox(E(x), E(y), E(w + 6), E(h + 4))
    tf = tb.text_frame
    tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE; tf.vertical_anchor = MSO_ANCHOR.TOP
    pl = t.get('pl') or 0; pt_ = t.get('pt') or 0
    tf.margin_left = E(pl); tf.margin_right = E(pl); tf.margin_top = E(pt_); tf.margin_bottom = 0
    # 依 <br> 分段
    paras = [[]]
    for r in t['runs']:
        if r.get('br'): paras.append([])
        else: paras[-1].append(r)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = max(1.0, t['lh'] * 0.85)
        p.alignment = PP_ALIGN.CENTER if t['align'] == 'center' else PP_ALIGN.LEFT
        if t['li'] and i == 0:
            r0 = p.add_run(); r0.text = '•  '
            set_font(r0, FONT_BODY, sz, False, TERRA)
        for r in runs:
            if not r.get('t'): continue
            run = p.add_run(); run.text = r['t']
            head = 'WenKai' in (r.get('f') or '')
            set_font(run, FONT_HEAD if head else FONT_BODY, r.get('sz') or sz, r.get('b'), rgb(r.get('c')))
    return tb

data = json.load(open('pptx_build/layout.json', encoding='utf-8'))
prs = Presentation(); prs.slide_width = E(1280); prs.slide_height = E(720)
blank = prs.slide_layouts[6]
for sd in data:
    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    # 裝飾圓（同網頁左下角）
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, E(-140), E(720 - 240), E(420), E(420))
    c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xED, 0xB4, 0x5A); alpha(c, 12); c.line.fill.background(); c.shadow.inherit = False
    if sd['cover']:
        for (cx, cy, r, col, a) in [(180, 60, 330, (0xFD, 0xE9, 0xD4), 70), (880, 650, 300, (0xE6, 0xEF, 0xE6), 70)]:
            o = s.shapes.add_shape(MSO_SHAPE.OVAL, E(cx - r), E(cy - r), E(2 * r), E(2 * r))
            o.fill.solid(); o.fill.fore_color.rgb = RGBColor(*col); alpha(o, a); o.line.fill.background(); o.shadow.inherit = False
    # 卡片
    for b in sd['boxes']:
        fill = rgb(b['bg']); line = rgb(b['line']); left = rgb(b['left'])
        if b['leftW'] > 2 and left and b['lineStyle'] != 'dashed':
            rrect(s, b['x'], b['y'], b['w'], b['h'], left, b['radius'])
            rrect(s, b['x'] + b['leftW'], b['y'], b['w'] - b['leftW'], b['h'], fill or BG, b['radius'], line, b['lineW'])
        else:
            rrect(s, b['x'], b['y'], b['w'], b['h'], fill, b['radius'], line, b['lineW'], b['lineStyle'] == 'dashed')
    # 文字
    for t in sd['texts']:
        textbox(s, t)
    # 頁碼
    if not sd['cover']:
        textbox(s, {'x': 56, 'y': 692, 'w': 120, 'h': 18, 'sz': 13, 'li': False, 'lh': 1.2, 'align': 'left',
                    'runs': [{'t': sd['num'], 'c': 'rgb(140,117,99)', 'sz': 13}]})
    # 影片框＋影片
    if sd['clip'] and sd['clipRect']:
        r = sd['clipRect']
        rrect(s, r['x'], r['y'], r['w'], r['h'], RGBColor(0xFF, 0xF8, 0xEE), 28, RGBColor(0xFF, 0xFF, 0xFF), 6)
        inset = 10; iw = r['w'] - 2 * inset; ih = iw * 16 / 9
        if ih > r['h'] - 2 * inset: ih = r['h'] - 2 * inset; iw = ih * 9 / 16
        x = r['x'] + (r['w'] - iw) / 2; y = r['y'] + (r['h'] - ih) / 2
        mv = s.shapes.add_movie(f"video/{sd['clip']}.mp4", E(x), E(y), E(iw), E(ih),
                                poster_frame_image=f"video/{sd['clip']}.jpg", mime_type='video/mp4')
        el = s.element
        for tm in el.findall(f'{{{NS}}}timing'): el.remove(tm)
        el.append(etree.fromstring(timing(mv.shape_id)))
    s.notes_slide.notes_text_frame.text = sd['note'].strip()
prs.save('個案研討簡報.pptx')
print('slides', len(data))
